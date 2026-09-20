import os

IMAGE_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.webp', '.gif', '.bmp', '.tif', '.tiff'}

def is_image_file(path):
    return os.path.splitext(path)[1].lower() in IMAGE_EXTENSIONS

def validate_image(path):
    """
    Integrity check for photos/scans of paper documents. Returns
    (ok, reason). A decodable image is NOT 'unreadable' — the vision
    model handles transcription downstream.
    """
    try:
        from PIL import Image
        with Image.open(path) as im:
            im.verify()
        return True, None
    except Exception as e:
        return False, f"Undecodable image file ({e})"

def parse_txt(path):
    with open(path, 'r', encoding='utf-8', errors='replace') as f:
        return f.read()

def parse_pdf(path):
    try:
        import pdfplumber
        text_content = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                text = page.extract_text(layout=True)
                if text:
                    text_content.append(text)
        return "\n".join(text_content)
    except Exception as e:
        print(f"Error parsing PDF {path}: {e}")
        return ""

def parse_docx(path):
    try:
        import docx
        doc = docx.Document(path)
        content = []
        for p in doc.paragraphs:
            if p.text.strip():
                content.append(p.text.strip())
        
        for table in doc.tables:
            for row in table.rows:
                row_data = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                content.append(" | ".join(row_data))
        return "\n".join(content)
    except Exception as e:
        print(f"Error parsing DOCX {path}: {e}")
        return ""

def parse_xlsx(path):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        content = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            for row in ws.iter_rows(values_only=True):
                if any(c is not None for c in row):
                    row_data = [str(c).strip().replace("\n", " ") if c is not None else "" for c in row]
                    content.append(" | ".join(row_data))
        return "\n".join(content)
    except Exception as e:
        print(f"Error parsing XLSX {path}: {e}")
        return ""

def extract_text(path):
    if not os.path.exists(path):
        return ""
    
    if os.path.getsize(path) == 0:
        return ""
        
    ext = os.path.splitext(path)[1].lower()

    # Photos/scans carry no embedded text — they are read by the vision
    # model (analyze_document_image), never by the text parsers.
    if ext in IMAGE_EXTENSIONS:
        return ""

    if ext == '.txt':
        return parse_txt(path)
    elif ext == '.pdf':
        return parse_pdf(path)
    elif ext == '.docx':
        return parse_docx(path)
    elif ext == '.xlsx':
        return parse_xlsx(path)
    else:
        return parse_txt(path)


# ---------------------------------------------------------------------------
# Tier 1: Fast Deterministic Field Extractor
# Extracts the 7 core fields from layout-parsed text (.txt, .pdf, .docx, .xlsx)
# ---------------------------------------------------------------------------
# An optional header qualifier is restricted to a parenthetical "(KG)" or a
# short "/exporter"-style suffix — an arbitrary greedy suffix can swallow the
# field value itself (e.g. "Shipper        APRIL FINE PAPER TRADING").
_PAREN = r'(?:\s*\([^)]*\))*'
# Field separator: colon/pipe, a run of 2+ spaces, or a single space when it
# directly follows a parenthetical qualifier (e.g. "(POD) FREMANTLE").
_FSEP = r'(?:[:|]|[ \t]{2,}|(?<=\))[ \t])'

FIELD_PATTERNS = {
    'shipper': [
        rf'(?:^|\|)[ \t]*shipper(?:\s*/\s*exporter)?{_PAREN}[ \t]*{_FSEP}[ \t]*([^\n|]*(?:\s*\|\s*on\s*behalf\s*of\s*[^|\n;]+)?)',
    ],
    'consignee': [
        rf'(?:^|\|)[ \t]*to\s*the\s*order\s*of{_PAREN}[ \t]+([^\n|]+)',
        rf'(?:^|\|)[ \t]*(?:consignee|to\s*order\s*of){_PAREN}[ \t]*{_FSEP}[ \t]*([^\n|]*)',
    ],
    'notify_party': [
        rf'(?:^|\|)[ \t]*notify(?:\s*party)?{_PAREN}[ \t]*{_FSEP}[ \t]*([^\n|]*)',
    ],
    'port_of_loading': [
        rf'(?:^|\|)[ \t]*(?:port\s*of\s*loading|load\s*port|\bpol\b){_PAREN}[ \t]*{_FSEP}[ \t]*([^\n|]+)',
        r'(?:^|\|)[ \t]*(?:port\s*of\s*loading|load\s*port)[ \t]+([A-Z][^\n|]+)',
    ],
    'port_of_discharge': [
        rf'(?:^|\|)[ \t]*(?:port\s*of\s*discharge|discharge\s*port|\bpod\b){_PAREN}[ \t]*{_FSEP}[ \t]*([^\n|]+)',
        r'(?:^|\|)[ \t]*(?:port\s*of\s*discharge|discharge\s*port)[ \t]+([A-Z][^\n|]+)',
    ],
    'container_count': [
        rf'(?:^|\|)[ \t]*(?:total\s*(?:no\.?\s*of\s*)?containers?|no\.?\s*of\s*containers?|container\s*count|containers?)(?!\s*no\b)(?:\s+or\s+packages)?{_PAREN}[ \t]*{_FSEP}[ \t]*([^\n|]+)',
        r'(?:^|\|)[ \t]*total\s*containers?[ \t]*[:|][ \t]*([^\n|]+)',
        r'(?:^|\|)[ \t]*containers?[ \t]*[:|][ \t]*([^\n|]+)',
    ],
    'gross_weight_kg': [
        rf'(?:^|\|)[ \t]*(?:(?:total\s+)?gross\s*(?:weight|wt)|weight){_PAREN}[ \t]*{_FSEP}[ \t]*([^\n|]*)',
    ]
}

def extract_shipping_fields_fast(text):
    """
    Rapid deterministic extraction of the 7 core fields.
    Returns a dict with all 7 fields populated (or empty strings if missing).
    Supports both inline 'Header: Value' and multi-line 'Header:\nValue'.
    """
    import re
    res = {}
    for fld, pat_list in FIELD_PATTERNS.items():
        val = ''
        for pat in pat_list:
            m = re.search(pat, text, re.IGNORECASE | re.MULTILINE)
            if m:
                matched_val = m.group(1).strip()
                if '|' in matched_val:
                    parts = [p.strip() for p in matched_val.split('|') if p.strip()]
                    matched_val = " ".join(parts)
                if matched_val:
                    val = matched_val
                    break

        # Fallback: if value is on the subsequent line (e.g. "Shipper:\nACME CORP")
        if not val:
            header_keys = {
                'shipper': r'(?:shipper(?:/exporter)?|exporter)',
                'consignee': r'(?:consignee(?:\s*\(non-negotiable\))?|to\s*the\s*order\s*of|to\s*order\s*of)',
                'notify_party': r'(?:notify\s*party|notify)',
                'port_of_loading': r'(?:port\s*of\s*loading|load\s*port|\bpol\b)',
                'port_of_discharge': r'(?:port\s*of\s*discharge|discharge\s*port|\bpod\b)',
                'container_count': r'(?:total\s*containers?|no\.?\s*of\s*containers?|container\s*count|containers?)',
                'gross_weight_kg': r'(?:(?:total\s+)?gross\s*weight|(?:total\s+)?gross\s*wt|weight)'
            }
            if fld in header_keys:
                hdr = header_keys[fld]
                next_line_m = re.search(rf'(?:^|\|)[ \t]*{hdr}[ \t]*[:|]?[ \t]*\n+[ \t]*([^\n|]+)', text, re.IGNORECASE | re.MULTILINE)
                if next_line_m:
                    candidate = next_line_m.group(1).strip()
                    # Don't pick up the next section header as the value
                    if not any(re.match(rf'^{h}[ \t]*[:|]?', candidate, re.IGNORECASE) for h in header_keys.values()):
                        val = candidate

        res[fld] = val
    return res
